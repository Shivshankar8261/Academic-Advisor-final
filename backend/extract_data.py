"""
Step 1 of the data pipeline: convert the university's RAW documents (data/raw/)
into clean, structured, citable files (data/processed/).

    raw PDF / XLSX  ──►  extract_data.py  ──►  processed .md / .json  ──►  ingest.py  ──►  ChromaDB

Nothing here is synthetic: every fact comes from the four files supplied with the
assignment. The only additions are clearly labelled:
  * Table 2 (letter grades) is an IMAGE inside the handbook PDF, so it is transcribed by hand.
  * term_context.md states how "current semester per batch" is derived (an assumption).

Run:  python extract_data.py
"""
import json
import re
import sqlite3
from collections import defaultdict
from pathlib import Path

import openpyxl
from pypdf import PdfReader

from config import COURSES_DB, PROCESSED_DIR, RAW_DIR

HANDBOOK_PDF = RAW_DIR / "student_handbook_aug2026.pdf"
SOP_PDF = RAW_DIR / "sop_students_aug2026.pdf"
SPREAD_XLSX = RAW_DIR / "semester_spread_structures_sept2026.xlsx"
MINOR_XLSX = RAW_DIR / "minor_courses_btech.xlsx"

# ---------------------------------------------------------------------------
# Handbook (PDF) -> handbook.md with one heading per section / clause
# ---------------------------------------------------------------------------

# (title, first_pdf_page, last_pdf_page) — found by inspecting the PDF page by page.
HANDBOOK_SECTIONS = [
    ("Section I - Preamble, Definitions", 9, 11),
    ("Section II - Admission Rules, Fee Policy and Scholarship Policy", 15, 20),
    ("Section III - Academic Regulations", 23, 47),
    ("Section IV - Code of Conduct", 51, 54),
    ("Section V - Anti-Ragging Guidelines", 57, 61),
    ("Section VI - Prevention of Sexual Harassment", 66, 68),
    ("Section VII - Ombudsperson", 71, 71),
    ("Section VIII - Dress Code", 75, 76),
    ("Section IX - Malpractice in Examination Hall", 79, 80),
    ("Section X - Knowledge Resource Centre (Library)", 83, 83),
    ("Section XI - Sports Policy", 87, 90),
]

# The last pages of the handbook, kept as their own section (TOC entry XII).
CONTACTS_PAGE = 91

# PDF pages 36 and 65 are two layered pages: each contains BOTH the text of printed page 28 (clauses 8.14-8.16)
# and printed page 57 (start of the POSH guidelines). The POSH block is removed from Section III and put back at
# the start of Section VI, where it belongs.
POSH_BLOCK_RE = re.compile(r"Guidelines on Prevention of Sexual Harassment\n1\. Introduction.*?"
                           r"hostile learning environment;\n", re.S)

# Table 1 is a bordered table whose cells pypdf reads as one word-wrapped column; re-typed from PDF page 31.
EVAL_TABLE_MD = """
Table 1. Evaluation System: Components and Weightage (re-typed from handbook page 23):

| Type of Course Structure | Evaluation Components | Weightage |
|---|---|---|
| Lecture-based Course (L in the L-T-P structure predominates, e.g. 3-0-0, 2-1-0, 3-0-2) | Continuous Assessments | 60% (minimum) to 70% (maximum) |
| Lecture-based Course | End Term Examination / Comprehensive | 30% (minimum) to 40% (maximum) |
| Practice-based Course (P in the L-T-P structure predominates, e.g. 0-0-4, 1-0-4, 1-0-2) | Continuous Assessments | 70% (minimum) to 100% (maximum) |
| Practice-based Course | End Term Examination / Jury / Project / Viva Voce | 30% (maximum) |
| Practice/Skill based courses (Industry Internship, Capstone project, Research Dissertation, Integrative Studio, Interdisciplinary Project, Seminar, Summer / Short Internship, Social Engagement / Field Projects and similar courses without a typical L-T-P structure; refer Clause 3.2) | Guidelines for the components of evaluation, with recommended weightages / criteria, are specified in the concerned Program Regulations and Course Plans | as applicable |
"""

# Table 3 is also a bordered table; text extraction scrambles which CGPA belongs to which year. Re-typed from
# PDF page 42 (printed page 34).
PROGRESSION_TABLE_MD = """
Table 3: Progression Criteria (re-typed from handbook page 34):

| Progression Stage | Minimum CGPA Requirement |
|---|---|
| Progression to Year 2 of the Program (i.e. to register for Semester 3) | Minimum CGPA of 4.00 (Eligibility Criteria to Register for Semester 3) |
| Progression to Year 3 and higher years of the Program (i.e. to register for Semester 5 and above) | Minimum CGPA of 5.00 |
"""

# (first line of the flattened table, clean markdown, regex of the line where the table ends)
FLATTENED_TABLES = [
    ("Table 1. Evaluation System", EVAL_TABLE_MD, r"^8\.8\."),
    ("Table 3: Progression Criteria", PROGRESSION_TABLE_MD, r"^12\.2\."),
]

# Table 2 of the handbook is an embedded image; transcribed manually from PDF page 32.
GRADE_TABLE_MD = """
Table 2. Letter Grades with Grade Points and Brief Qualitative Description (transcribed from the image on handbook page 24):

| Letter Grade | Grade Point | Qualitative Description |
|---|---|---|
| O | 10 | Outstanding |
| A+ | 9 | Excellent |
| A | 8 | Very Good |
| B+ | 7 | Good |
| B | 6 | Above Average |
| C | 5 | Average |
| D | 4 | Pass |
| F | 0 | Fail |
| FA | 0 | Fail - Shortage of Attendance |
| S | - | Satisfactory |
| U | - | Unsatisfactory |
| I | - | Incomplete |
"""

# Top-level clause titles of the Academic Regulations (clause 1 has no number in the PDF).
ACADEMIC_CLAUSES = {
    1: "Academic Calendar", 2: "Registration", 3: "Academic Credits - Course Credit Structure",
    4: "Program Regulations", 5: "Medium of Instruction and Evaluation",
    6: "Maximum Duration for the Completion of a Program", 7: "Attendance Requirements",
    8: "Teaching, Evaluation and Grading System", 9: "Examinations", 10: "Academic Appeals",
    11: "Summer Term", 12: "Progression", 13: "Withdrawal/Re-joining", 14: "Transfer of Credits",
    15: "Award of Degree", 16: "Convocation", 17: "Power To Revise, Modify and Amend",
}


def clean_lines(text: str) -> list[str]:
    out = []
    for line in text.splitlines():
        line = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", " ", line)  # PDF control chars (e.g. \x03)
        line = re.sub(r"\s+", " ", line).strip()
        if not line or re.fullmatch(r"\d{1,3}", line):  # drop blank lines & bare page numbers
            continue
        out.append(line)
    return out


def academic_regulations_md(lines: list[str]) -> str:
    """Split Section III into '### N. Title' headings using the known clause titles."""
    md, current, skip_until = [], 0, None
    for line in lines:
        table = next((t for t in FLATTENED_TABLES if line.startswith(t[0])), None)
        if table:
            md.append(table[1])
            skip_until = table[2]
            continue
        if skip_until:  # drop the word-wrapped cells of the table; it ends where the next clause starts
            if not re.match(skip_until, line):
                if line.startswith("<!-- page:"):
                    md.append(line)
                continue
            skip_until = None
        for num, title in ACADEMIC_CLAUSES.items():
            key = title.split()[0]
            # e.g. "2. Registration", "12.Progression", "10.Academic Appeals", or bare "Academic Calendar"
            if num > current and re.match(rf"^({num}\.\s*)?{re.escape(key)}", line) and len(line) < 70 \
                    and (num == 1 or line.startswith(f"{num}.")):
                md.append(f"\n### {num}. {title}\n")
                current = num
                line = ""
                break
        if line:
            md.append(line)
            if "summarized in Table 2" in line:
                md.append(GRADE_TABLE_MD)
    return "\n".join(md)


HANDBOOK_PAGE_OFFSET = 8   # PDF page 32 carries the printed page number 24
ADDRESS_PAGE = 92


def page_marker(printed_page: int) -> str:
    """Invisible marker kept in the processed markdown; ingest.py turns it into page metadata for citations."""
    return f"<!-- page:{printed_page} -->"


def page_text(reader: PdfReader, p: int, marker: bool = True) -> str:
    lines = clean_lines(reader.pages[p - 1].extract_text() or "")
    head = [page_marker(p - HANDBOOK_PAGE_OFFSET)] if marker else []
    return "\n".join(head + lines) + "\n"


def build_handbook_md() -> str:
    reader = PdfReader(str(HANDBOOK_PDF))
    parts = ["# Vidyashilp University Student Handbook (August 2026)\n"]
    for title, start, end in HANDBOOK_SECTIONS:
        text = "".join(page_text(reader, p) for p in range(start, end + 1))
        if title.startswith("Section III"):
            text = POSH_BLOCK_RE.sub("", text)             # layered page 36 -> keep only printed page 28
        if title.startswith("Section VI"):                 # layered page 65 -> keep only printed page 57
            posh = POSH_BLOCK_RE.search(page_text(reader, 65, marker=False))
            assert posh, "POSH introduction not found on PDF page 65"
            text = page_marker(65 - HANDBOOK_PAGE_OFFSET) + "\n" + posh.group(0) + text
        lines = text.splitlines()
        parts.append(f"\n## {title}\n")
        if title.startswith("Section III"):
            parts.append("Amended Academic Regulations approved in the 6th Meeting of the Academic Council "
                         "held on July 29, 2025 vide Resolution No 6.16.")
            parts.append(academic_regulations_md(lines))
        else:
            parts.append("\n".join(lines))
    # "3 Deputy Registrar (Examination &" / "Evaluation)" / "7204176753" -> one "- role: number" line each
    raw = " ".join(line for line in clean_lines(reader.pages[CONTACTS_PAGE - 1].extract_text() or "")
                   if "ŵ" not in line)  # the page heading is in a font pypdf cannot decode
    contacts = [f"- {role.strip()}: {phone}" for role, phone in re.findall(r"\d\s+(.+?)\s+(\d{10})", raw)]
    parts.append("\n## Section XII - Important Contact Numbers\n")
    parts.append(page_marker(CONTACTS_PAGE - HANDBOOK_PAGE_OFFSET))
    parts.append("Phone numbers of university offices / officials students can contact:")
    parts.append("\n".join(contacts))
    # back cover: the university's postal address (the only place the documents state where VU is)
    address = " ".join(clean_lines(reader.pages[ADDRESS_PAGE - 1].extract_text() or ""))
    parts.append("\n## University Address (handbook back cover)\n")
    parts.append(page_marker(ADDRESS_PAGE - HANDBOOK_PAGE_OFFSET))
    parts.append(f"Location / postal address of Vidyashilp University: {address}")
    text = "\n".join(parts)
    # guard against the layered-page problem coming back
    assert text.count("8.15 Transcript") == 1 and "Vishaka" not in text.split("## Section IV")[0]
    return text


# ---------------------------------------------------------------------------
# SOP (PDF) -> sop.md, one heading per numbered SOP topic
# ---------------------------------------------------------------------------
SOP_TOPICS = ["Course Registration", "Student Attendance", "Change of Program", "Summer Term Registration",
              "University Annual Fee", "Digii Account", "Bonafide Certificate", "Scholarship", "Grievance",
              "General Guidelines", "Vehicle and Parking Guidelines"]
# First bullet of each topic (used as an anchor because the PDF puts topic labels in a side column).
SOP_ANCHORS = ["Eligibility Criteria", "Course wise attendance can be check", "Student may seek transfer",
               "The Summer Term is an additional", "The student shall remit the Annual", "Digii is a comprehensive",
               "Educational Loan", "Students seeking support or information regarding scholarships",
               "Students can submit their grievances", "To ensure a seamless university experience",
               "Vehicle access limitations"]


def build_sop_md() -> str:
    reader = PdfReader(str(SOP_PDF))
    lines = []
    for pno, page in enumerate(reader.pages, start=1):
        lines.append(page_marker(pno))
        for line in clean_lines(page.extract_text() or ""):
            if re.fullmatch(r"Page \d of \d", line):
                continue
            # remove side-column topic labels like "1. Course", "Registration", "2. Student", ...
            if re.fullmatch(r"\d{1,2}\.\s*[A-Z][A-Za-z ]{0,25}", line) or line in {
                    "Registration", "Attendance", "Program", "Annual Fee", "Account", "Certificate",
                    "Guidelines", "Parking", "Parking Guidelines"}:
                continue
            lines.append(line)
    text = "\n".join(lines)
    md = ["# Standard Operating Procedure (SOP) for Students - Vidyashilp University (17 Aug 2026)\n"]
    positions = [(text.find(a), i) for i, a in enumerate(SOP_ANCHORS)]
    positions = sorted(p for p in positions if p[0] >= 0)
    for k, (pos, i) in enumerate(positions):
        end = positions[k + 1][0] if k + 1 < len(positions) else len(text)
        md.append(f"\n## {i + 1}. {SOP_TOPICS[i]}\n")
        md.append(text[pos:end].strip())
    return "\n".join(md)


# ---------------------------------------------------------------------------
# Semester spread (XLSX) -> course catalogue, programme structure, semester offerings
# ---------------------------------------------------------------------------
SPREAD_SHEETS = {  # sheet name -> batch (year of admission)
    "Sem Spread BTECH-2022": "2022", "Sem_Spread_2023": "2023", "Sem_Spread_2024": "2024",
    "Sem_Spread_2025": "2025", "Sem_Spread_DS_2026": "2026",
}
# 1-based column of "Course Code" for each semester block; the block is code, title, prereq, L, T, P, C
SEM_CODE_COLS = {1: 5, 2: 12, 3: 20, 4: 27, 5: 35, 6: 42, 7: 50, 8: 57}
CODE_RE = re.compile(r"\b[A-Z]{3,4}\s?\d{3}\b")


def _s(v) -> str:
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def _num(v):
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return None


NIL_VALUES = {"", "NIL", "NONE", "-", "NA", "N/A"}
UNDECIDED_VALUES = {"TBA", "TBD"}


def parse_prereqs(text: str) -> list[str]:
    if not text or text.strip().upper() in NIL_VALUES | UNDECIDED_VALUES:
        return []
    return [c.replace(" ", "") for c in CODE_RE.findall(text.upper())] or [text.strip()]


def prereq_groups(text: str) -> list[list[str]]:
    """'ECON201, ECON202/ ECON207' -> [['ECON201'], ['ECON202', 'ECON207']]  (AND of OR-groups).
    In the source files a comma separates required courses and '/' separates alternatives."""
    if not text or text.strip().upper() in NIL_VALUES | UNDECIDED_VALUES:
        return []
    groups = []
    for part in re.split(r",", text.upper()):
        codes = [c.replace(" ", "") for c in CODE_RE.findall(part)]
        if codes:
            groups.append(list(dict.fromkeys(codes)))
    return groups


def prereq_text(raw: str) -> str:
    """Normalised, human-readable prerequisite rule."""
    r = (raw or "").strip()
    if r.upper() in NIL_VALUES:
        return "none"
    if r.upper() in UNDECIDED_VALUES:
        return "not yet decided (TBA/TBD in source)"
    groups = prereq_groups(r)
    if not groups:  # no course code at all, e.g. "FAMA Financial and management accounting"
        return f"'{r}' (no course code given in the source)"
    return " AND ".join(g[0] if len(g) == 1 else "(" + " OR ".join(g) + ")" for g in groups)


def parse_spreads():
    wb = openpyxl.load_workbook(SPREAD_XLSX, data_only=True)
    offerings = []   # one record per (batch, semester, course slot)
    baskets = {}     # batch -> {basket: (fixed minimum credits, credits actually placed in the spread)}
    for sheet, batch in SPREAD_SHEETS.items():
        ws = wb[sheet]
        basket = ""
        baskets[batch] = {}
        for rnum, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            row = list(row) + [None] * 70
            label = _s(row[1])
            if label and _num(row[3]) is not None:              # new basket row e.g. "B3 | Program CORE | 46"
                basket = label
                # col D "Fixed Min", col BM "Total Credits" (formula, cached value), sheet row
                baskets[batch][basket] = (_num(row[3]), _num(row[64]), rnum)
            if label.lower().startswith("total") or _s(row[0]).lower().startswith("total"):
                break
            for sem, c in SEM_CODE_COLS.items():
                code, title, pre = _s(row[c - 1]), _s(row[c]), _s(row[c + 1])
                L, T, P, C = (_num(row[c + k]) for k in (2, 3, 4, 5))
                if not code and not title:
                    continue
                above = next((o for o in reversed(offerings) if o["batch"] == batch and o["semester"] == sem), None)
                if not title and C is None and above and code.replace(" ", "") in above["course_code"].replace(" ", ""):
                    continue  # spill-over line of the cell above (e.g. 2022 S5 "MATH401/ COMP401" + "COMP401")
                if not title and not CODE_RE.search(code):       # placeholder slot, e.g. "Minor/Open (1 Course)"
                    # layout for placeholder rows: [label, credits, 0, 0, credits]
                    title, code, pre = code, "", ""
                    C = _num(row[c + 3]) or _num(row[c])
                    L = T = P = None
                offerings.append({
                    "batch": batch, "semester": sem, "basket": basket,
                    "course_code": re.sub(r"([A-Z]{3,4}) (\d{3})", r"\1\2", code) or "TBA", "title": title, "prerequisites_raw": pre or "NIL",
                    "prerequisites": parse_prereqs(pre), "L": L, "T": T, "P": P, "credits": C,
                    "source_tab": sheet, "source_row": rnum,
                })
    return offerings, baskets


def parse_minors():
    wb = openpyxl.load_workbook(MINOR_XLSX, data_only=True)
    records = []
    for ws in wb.worksheets:
        minor = ws.title.replace(" Minor", "").strip()
        batch = ""
        header = None
        for rnum, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = [_s(v) for v in row]
            joined = " ".join(vals)
            m = re.search(r"(20\d\d) Batch", joined)
            if m:
                batch = m.group(1)
            if "Course Code" in vals or "Course Title" in vals:
                header = vals
                continue
            if "No Students" in joined:
                records.append({"minor": minor, "batch": batch, "course_code": "-", "title":
                                "No students / minor not offered for this batch", "credits": None,
                                "semester": None, "prerequisites_raw": "", "prerequisites": [],
                                "source_tab": ws.title, "source_row": rnum})
                continue
            # find a course row: contains a title-like cell and a numeric credit & semester
            nums = [(i, _num(v)) for i, v in enumerate(vals) if _num(v) is not None]
            texts = [v for v in vals if v and _num(v) is None and not re.search(r"20\d\d Batch", v)
                     and v not in {"Finance Minor", "Marketing Minor", "Start -up Minor"}]
            if not texts:
                continue
            if texts[0].upper() in {"TBA", "TBD", "DON’T KNOW", "DON'T KNOW"} and nums and len(texts) == 1:
                records.append({"minor": minor, "batch": batch, "course_code": texts[0].upper(),
                                "title": "Courses not yet decided (TBA/TBD in source)", "credits": nums[-1][1],
                                "semester": None, "prerequisites_raw": "", "prerequisites": [],
                                "source_tab": ws.title, "source_row": rnum})
                continue
            code_idx = next((i for i, v in enumerate(vals) if CODE_RE.fullmatch(v.replace(" ", "")[:7]) or
                             v.upper() in {"NEW", "DON’T KNOW", "DON'T KNOW"} or "/" in v and CODE_RE.search(v)), None)
            if code_idx is None or len(nums) < 2:
                continue
            code = vals[code_idx]
            title = vals[code_idx + 1]
            after = [n for i, n in nums if i > code_idx + 1]
            # layout: L T P Credit Semester [Pre-rq] ; some "New" rows only have Credit Semester
            credit, sem = (after[3], after[4]) if len(after) >= 5 else (after[-2], after[-1])
            sem_idx = [i for i, n in nums if i > code_idx + 1][4 if len(after) >= 5 else -1]
            pre = " ".join(v for v in vals[sem_idx + 1:] if v)
            unknown = {"NEW": "NEW (code not assigned)", "DON’T KNOW": "Code not known (DON'T KNOW in source)",
                       "DON'T KNOW": "Code not known (DON'T KNOW in source)"}
            records.append({"minor": minor, "batch": batch, "course_code": unknown.get(code.upper(), code.upper()), "title": title, "credits": credit, "semester": sem,
                            "prerequisites_raw": pre or "Nil", "prerequisites": parse_prereqs(pre),
                            "source_tab": ws.title, "source_row": rnum})
    return records


LTP_RE = re.compile(r"\s*LTP:\s*(\d)-(\d)-(\d)")


def split_choice_slot(o) -> list[dict]:
    """Either/or slots like 'PSYC101/ECON101 | Introduction to Psychology/Introduction to Economics' or
    'MATH401/ COMP401 | SPT-Course #1-Probabilistic Graph Models LTP: 3-0-2 /Web Framework LTP: 2-0-4'
    -> one record per alternative (only when codes, titles and prerequisites line up)."""
    codes = [c.strip() for c in o["course_code"].split("/")]
    if len(codes) < 2:
        return []
    title = re.sub(r"^SPT-Course\s*#\d+\s*[-:]?\s*", "", o["title"]).strip()
    title = re.sub(r"^\((.*)\)$", r"\1", title)
    titles = [t.strip() for t in title.split("/")]
    pres = [p.strip() for p in o["prerequisites_raw"].split("/")]
    if len(titles) != len(codes):
        return []
    if len(pres) != len(codes):
        pres = [o["prerequisites_raw"]] * len(codes)
    slot = o["title"] if o["title"].startswith("SPT") else " / ".join(codes)
    out = []
    for code, t, pre in zip(codes, titles, pres):
        code = code.replace(" ", "").upper()
        if not CODE_RE.fullmatch(code):
            continue  # "TBA"
        ltp = LTP_RE.search(t)
        rec = dict(o, course_code=code, title=LTP_RE.sub("", t).strip(), prerequisites_raw=pre or "NIL",
                   prerequisites=parse_prereqs(pre), choice_of=f"either/or slot: {slot}")
        if ltp:
            rec.update(L=int(ltp.group(1)), T=int(ltp.group(2)), P=int(ltp.group(3)))
        out.append(rec)
    return out


def build_catalogue(offerings, minors):
    """Aggregate per course code -> the 'course catalogue'.

    Curricula differ by batch (e.g. Machine Learning's prerequisites changed for the 2025 batch),
    so every batch-specific detail is kept under `by_batch` instead of being collapsed.
    """
    cat = {}
    rows = offerings + [dict(m, basket=f"Minor - {m['minor']}") for m in minors]
    rows += [alt for o in rows for alt in split_choice_slot(o)]
    for o in sorted(rows, key=lambda r: r["batch"]):          # later batches overwrite headline fields
        code = o["course_code"]
        if not CODE_RE.fullmatch(code):
            continue  # skip placeholders (TBA) & either/or track slots like "MATH401/ COMP401" (kept in offerings)
        c = cat.setdefault(code, {"course_code": code, "by_batch": {}})
        c.update(title=o["title"], credits=o["credits"], basket=o["basket"])
        if o.get("L") is not None:
            c["L-T-P"] = f"{o['L']}-{o['T']}-{o['P']}"
        c["by_batch"][o["batch"]] = {
            "semester": o["semester"], "prerequisites": o["prerequisites"],
            "prerequisites_raw": o["prerequisites_raw"], "credits": o["credits"], "basket": o["basket"]}
        if o.get("choice_of"):
            c["by_batch"][o["batch"]]["choice_of"] = o["choice_of"]
    known = set(cat)
    for c in cat.values():
        prereqs = {p for b in c["by_batch"].values() for p in b["prerequisites"]}
        c["prerequisites_differ_by_batch"] = len({tuple(b["prerequisites"]) for b in c["by_batch"].values()}) > 1
        c["unresolved_prerequisites"] = sorted(p for p in prereqs if CODE_RE.fullmatch(p) and p not in known)
    return dict(sorted(cat.items()))


def programme_structure_md(offerings, baskets) -> str:
    md = ["# B.Tech Programme Structure and Semester Spread (source: Semester_Spread_Structures_Sept_2026.xlsx)\n",
          "Each batch (year of admission) has its own curriculum. Total credits for the degree: 180 for every batch."]
    by_batch = defaultdict(list)
    for o in offerings:
        by_batch[o["batch"]].append(o)
    for batch in sorted(by_batch):
        tab = next(t for t, b in SPREAD_SHEETS.items() if b == batch)
        md.append(f"\n## {batch} Batch - Credit baskets (tab {tab})\n")
        for b, (cr, placed, _row) in baskets[batch].items():
            note = f" (NOTE: the semester spread places {placed} credits in this basket)" \
                if placed is not None and placed != cr else ""
            md.append(f"- {b}: minimum {cr} credits{note}")
        for sem in range(1, 9):
            rows = [o for o in by_batch[batch] if o["semester"] == sem]
            if not rows:
                continue
            md.append(f"\n## {batch} Batch - Semester {sem} (tab {rows[0]['source_tab']})\n")
            md.append(f"Semester {sem} is an {'odd' if sem % 2 else 'even'} semester. Courses:")
            for o in rows:
                md.append(f"- {o['course_code']} {o['title']} | basket: {o['basket']} | credits: {o['credits']} "
                          f"| L-T-P: {o['L']}-{o['T']}-{o['P']} | prerequisites: {prereq_text(o['prerequisites_raw'])} "
                          f"| row {o['source_row']}")
            md.append(f"Total listed credits in semester {sem}: {sum((o['credits'] or 0) for o in rows)}")
    return "\n".join(md)


def minors_md(minors) -> str:
    md = ["# Minor Courses for B.Tech Students (source: Minor_Courses_for_BTech_Students.xlsx)\n",
          "Minor/Open basket is 32 credits for 2025/2026 batches and 24 for 2022-2024 batches (see programme structure). "
          "Entries marked TBA/TBD/DON'T KNOW are not yet decided in the source file."]
    groups = defaultdict(list)
    for m in minors:
        groups[(m["minor"], m["batch"])].append(m)
    for (minor, batch), rows in groups.items():
        md.append(f"\n## {minor} Minor - {batch} Batch (tab {rows[0]['source_tab']}, rows "
                  f"{min(m['source_row'] for m in rows)}-{max(m['source_row'] for m in rows)})\n")
        for m in rows:
            md.append(f"- {m['course_code']} {m['title']} | credits: {m['credits']} | semester: {m['semester']} "
                      f"| prerequisites: {prereq_text(m['prerequisites_raw'])} | row {m['source_row']}")
    return "\n".join(md)


STRUCT_SHEETS = {"Struct_2022": "2022", "Struct_2023": "2023", "Struct_2024": "2024", "Struct_2025": "2025",
                 "Struct_2026_DS": "2026"}
BASKET_NAMES = ["University Core", "Foundation", "Program Core", "Program Honors", "Specialization Tracks",
                "Electives", "Minor/Open", "Internship/Capstone Project"]


def struct_summary_md() -> str:
    """The Struct_* sheets are the summary 'credit structure' tables. They are kept separate from the
    semester spread because the two sometimes disagree (e.g. 2026 batch University Core 18 vs 24)."""
    wb = openpyxl.load_workbook(SPREAD_XLSX, data_only=True)
    md = ["# Programme Credit Structure summary sheets (source: Semester_Spread_Structures_Sept_2026.xlsx, Struct_* sheets)\n",
          "These summary sheets list the credit baskets and course lists per batch. Where they differ from the "
          "semester-spread sheets of the same file, both values are shown in the respective documents."]
    for sheet, batch in STRUCT_SHEETS.items():
        ws = wb[sheet]
        baskets, rows = {}, []
        for row in ws.iter_rows(values_only=True):
            vals = [_s(v) for v in row]
            for i, v in enumerate(vals):
                name = next((b for b in BASKET_NAMES if v.lower().replace(" ", "") == b.lower().replace(" ", "")), None)
                if name and name not in baskets:
                    nxt = next((_num(x) for x in vals[i + 1:] if x), None)
                    if nxt is not None:
                        baskets[name] = nxt
            if any(vals):
                rows.append(" | ".join(v for v in vals if v))
        md.append(f"\n## {batch} Batch - Credit structure summary ({sheet})\n")
        md += [f"- {b}: {c} credits" for b, c in baskets.items()]
        md.append(f"- Total credits for the degree: 180")
        md.append(f"\n## {batch} Batch - Course lists by basket ({sheet})\n")
        md.append("Rows of the summary table (columns: #, course, credits, grouped by basket headings):")
        md += rows
    return "\n".join(md)


SPREAD_FILE = "Semester_Spread_Structures_Sept_2026.xlsx"   # names cited to students (the original documents)
MINOR_FILE = "Minor_Courses_for_BTech_Students.xlsx"


def struct_baskets() -> list[dict]:
    """Basket credits as stated in the Struct_* summary sheets (compared against the spread sheets)."""
    wb = openpyxl.load_workbook(SPREAD_XLSX, data_only=True)
    out = []
    for sheet, batch in STRUCT_SHEETS.items():
        seen = set()
        for rnum, row in enumerate(wb[sheet].iter_rows(values_only=True), start=1):
            vals = [_s(v) for v in row]
            for i, v in enumerate(vals):
                name = next((b for b in BASKET_NAMES if v.lower().replace(" ", "") == b.lower().replace(" ", "")), None)
                nxt = next((_num(x) for x in vals[i + 1:] if x), None) if name else None
                if name and name not in seen and nxt is not None:
                    seen.add(name)
                    out.append({"batch": batch, "basket": name, "credits": nxt, "source_tab": sheet, "source_row": rnum})
    return out


def code_status(code: str) -> str:
    if code.startswith("NEW"):
        return "code_not_assigned"
    if code.startswith("Code not known"):
        return "code_unknown"
    if code in {"TBA", "TBD", "DON’T KNOW", "DON'T KNOW"}:
        return "placeholder"
    if code == "-":
        return "not_offered"
    if "/" in code:
        return "choice_slot"
    return "assigned" if CODE_RE.fullmatch(code) else "placeholder"


def build_store(offerings, baskets, minors) -> int:
    """Structured store for exact lookups (counts, lists, credits, prerequisites, semesters):
    one row per course slot, with the original file / tab / row it came from."""
    COURSES_DB.unlink(missing_ok=True)
    db = sqlite3.connect(COURSES_DB)
    db.executescript("""
    CREATE TABLE courses (id INTEGER PRIMARY KEY, kind TEXT, batch TEXT, semester INTEGER, basket TEXT, minor TEXT,
        course_code TEXT, code_status TEXT, title TEXT, L INTEGER, T INTEGER, P INTEGER, credits REAL,
        prerequisites_raw TEXT, prerequisites_normalized TEXT, prerequisite_codes TEXT, choice_of TEXT,
        source_file TEXT, source_tab TEXT, source_row INTEGER);
    CREATE TABLE baskets (batch TEXT, basket TEXT, min_credits REAL, placed_credits REAL,
        source_file TEXT, source_tab TEXT, source_row INTEGER);
    CREATE TABLE struct_baskets (batch TEXT, basket TEXT, credits REAL, source_file TEXT, source_tab TEXT,
        source_row INTEGER);
    CREATE INDEX idx_code ON courses(course_code);
    """)
    rows = []

    def add(o, kind, minor=None, choice_of=None):
        code = o["course_code"]
        rows.append((kind, o["batch"], o["semester"], o.get("basket"), minor, code,
                     "assigned" if choice_of else code_status(code), o["title"], o.get("L"), o.get("T"), o.get("P"),
                     o["credits"], o["prerequisites_raw"], prereq_text(o["prerequisites_raw"]),
                     ",".join(c for g in prereq_groups(o["prerequisites_raw"]) for c in g), choice_of,
                     SPREAD_FILE if kind == "programme" else MINOR_FILE, o["source_tab"], o["source_row"]))

    for o in offerings:
        add(o, "programme")
        for alt in split_choice_slot(o):
            add(alt, "programme", choice_of=alt["choice_of"])
    for m in minors:
        add(dict(m, basket=f"Minor - {m['minor']}"), "minor", minor=m["minor"])
        for alt in split_choice_slot(dict(m, basket=f"Minor - {m['minor']}")):
            add(alt, "minor", minor=m["minor"], choice_of=alt["choice_of"])
    db.executemany(f"INSERT INTO courses VALUES (NULL{', ?' * 19})", rows)
    tab_of = {b: t for t, b in SPREAD_SHEETS.items()}
    db.executemany("INSERT INTO baskets VALUES (?, ?, ?, ?, ?, ?, ?)",
                   [(batch, name, mn, placed, SPREAD_FILE, tab_of[batch], r)
                    for batch, bs in baskets.items() for name, (mn, placed, r) in bs.items()])
    db.executemany("INSERT INTO struct_baskets VALUES (?, ?, ?, ?, ?, ?)",
                   [(b["batch"], b["basket"], b["credits"], SPREAD_FILE, b["source_tab"], b["source_row"])
                    for b in struct_baskets()])
    db.commit()
    n = db.execute("SELECT COUNT(*) FROM courses").fetchone()[0]
    db.close()
    return n


TERM_CONTEXT_MD = """# Current Academic Term Context (derived - see note)

## How the current semester is determined
- The course-offering file is titled "Semester Spread Structures Sept 2026".
- Handbook clause 1.2: the Odd Semester normally runs July/August to December; the Even Semester runs January to April/May.
- Therefore the current term (September 2026) is the ODD semester of academic year 2026-27, and the NEXT semester is the EVEN semester (January-May 2027).
- A summer term (June-July 2027) may be offered, but which courses run in it is not specified (handbook 11.4, 11.5.5).

## Current semester of each batch in Sept 2026 (derived from batch year)
- 2026 batch: Semester 1 (next: Semester 2)
- 2025 batch: Semester 3 (next: Semester 4)
- 2024 batch: Semester 5 (next: Semester 6)
- 2023 batch: Semester 7 (next: Semester 8)
- 2022 batch: normal 4-year duration completed (maximum duration is N+2 = 6 years per handbook 6.1)

## Offering rule of thumb
Courses in the semester spread are placed in a specific semester. Odd-semester courses (S1, S3, S5, S7) run in the odd term and even-semester courses (S2, S4, S6, S8) run in the even term. A course placed in an even semester is therefore not offered in the current (odd) term. Offerings of specialization/minor courses also require at least 10 registered students unless the Program Chair permits otherwise (handbook 2.13).
"""


def main():
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    (PROCESSED_DIR / "handbook.md").write_text(build_handbook_md())
    (PROCESSED_DIR / "sop.md").write_text(build_sop_md())

    offerings, baskets = parse_spreads()
    minors = parse_minors()
    catalogue = build_catalogue(offerings, minors)

    (PROCESSED_DIR / "semester_offerings.json").write_text(json.dumps(offerings, indent=1))
    (PROCESSED_DIR / "course_catalogue.json").write_text(json.dumps(catalogue, indent=1))
    (PROCESSED_DIR / "minors.json").write_text(json.dumps(minors, indent=1))
    (PROCESSED_DIR / "programme_structure.md").write_text(programme_structure_md(offerings, baskets))
    (PROCESSED_DIR / "minors.md").write_text(minors_md(minors))
    (PROCESSED_DIR / "term_context.md").write_text(TERM_CONTEXT_MD)
    (PROCESSED_DIR / "structure_summary.md").write_text(struct_summary_md())
    n = build_store(offerings, baskets, minors)
    print(f"courses.db: {n} course rows (structured store for exact lookups)")
    print(f"handbook.md, sop.md written; {len(offerings)} offering slots, {len(catalogue)} catalogue courses, "
          f"{len(minors)} minor rows")


if __name__ == "__main__":
    main()
